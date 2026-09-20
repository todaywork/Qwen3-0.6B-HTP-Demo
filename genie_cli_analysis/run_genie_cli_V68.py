#!/usr/bin/env python3
"""Deploy Genie CLI and run single or batch inference.

Batch mode reads ``输入提示词`` and ``预期结果``. By default a resident
device process loads the model once, journals each result, and exports Excel
after the batch. --batch-engine legacy retains the per-process baseline.
"""

from __future__ import annotations

import argparse
import base64
import json
import re
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from tokenizers import Tokenizer


LOCAL_ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = LOCAL_ROOT / "smoke20_input.xlsx"
DEFAULT_REFERENCE = (
    LOCAL_ROOT / "batch_result_Qwen3-0.6B-Safetensors_20260516_131103.xlsx"
)
DEFAULT_DEVICE_ROOT = "/data/local/tmp/genie_qwen3_quality"
DEFAULT_TOKENIZER = LOCAL_ROOT.parent / "tokenizer.json"
ANDROID_CONTEXT_SIZE = 512
ANDROID_MAX_ALL_TOKENS = 256
ANDROID_MAX_OUTPUT_TOKENS = 48
ANDROID_BOS_TOKEN_COUNT = 1
ANDROID_THREAD_COUNT = 12
ANDROID_SEED = 42
ANDROID_TOP_K = 40
ANDROID_TOP_P = 0.95
ANDROID_TEMPERATURE = 0.0
ANDROID_PRESENCE_PENALTY = 0.0
HEADERS = [
    "输入提示词",
    "预期结果",
    "输入TokenIDs",
    "推理结果",
    "比对结果",
    "Prefill(s)",
    "Decode(s)",
    "Total(s)",
    "Prefill(tok/s)",
    "Decode(tok/s)",
    "PromptTokens",
    "DecodeTokens",
]
REQUIRED_FILES = [
    Path("bin/genie-t2t-run"),
    Path("libs/libGenie.so"),
    Path("libs/libQnnHtp.so"),
    Path("libs/libQnnHtpNetRunExtensions.so"),
    Path("libs/libQnnHtpV68Stub.so"),
    Path("libs/libQnnSystem.so"),
    Path("genie_cli_prompt.txt"),
]


@dataclass
class Metrics:
    prefill_seconds: float
    decode_seconds: float
    total_seconds: float
    prefill_rate: float
    decode_rate: float
    prompt_tokens: int
    decode_tokens: int


@dataclass
class Case:
    prompt: str
    expected: str


def configure_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8", errors="replace")


def run_adb(*args: str, capture: bool = False) -> str:
    command = ["adb", *args]
    completed = subprocess.run(
        command,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    output = completed.stdout or ""
    if completed.returncode != 0:
        raise RuntimeError(
            f"ADB command failed ({completed.returncode}): {' '.join(command)}\n{output}"
        )
    return output


def shell_base64(script: str, *, use_shell_user: bool = False) -> str:
    encoded = base64.b64encode(script.replace("\r\n", "\n").encode("utf-8")).decode(
        "ascii"
    )
    suffix = "sh 2>&1"
    if use_shell_user:
        uid = run_adb("shell", "id -u", capture=True).strip()
        if uid != "2000":
            if uid != "0":
                raise RuntimeError(f"Expected ADB shell or root identity, got uid={uid}")
            # Root-enabled engineering devices may need an explicit drop to shell.
            # Production phones already run adbd commands as shell and have no su.
            run_adb("shell", "command -v su", capture=True)
            suffix = "su shell sh 2>&1"
    return run_adb("shell", f"echo '{encoded}' | base64 -d | {suffix}", capture=True)


def elf_class(path: Path) -> str:
    header = path.read_bytes()[:5]
    if len(header) != 5 or header[:4] != b"\x7fELF":
        return "not-elf"
    return {1: "ELF32", 2: "ELF64"}.get(header[4], "unknown")


def validate_inputs(args: argparse.Namespace) -> None:
    if shutil.which("adb") is None:
        raise RuntimeError("adb was not found in PATH")
    if not re.fullmatch(r"/data/local/tmp/[A-Za-z0-9._/-]+", args.device_root):
        raise ValueError(f"Unsafe device root: {args.device_root}")
    if "/../" in args.device_root or args.device_root.endswith("/.."):
        raise ValueError(f"Unsafe device root: {args.device_root}")
    if not Path(args.tokenizer).is_file():
        raise FileNotFoundError(f"Tokenizer file does not exist: {args.tokenizer}")

    missing = [str(LOCAL_ROOT / path) for path in REQUIRED_FILES if not (LOCAL_ROOT / path).is_file()]
    if missing:
        raise FileNotFoundError("Missing required local files:\n" + "\n".join(missing))
    invalid = [
        str(LOCAL_ROOT / path)
        for path in REQUIRED_FILES
        if path.name != "genie_cli_prompt.txt" and elf_class(LOCAL_ROOT / path) != "ELF64"
    ]
    if invalid:
        raise RuntimeError("Android host binaries are not ELF64:\n" + "\n".join(invalid))

    state = run_adb("get-state", capture=True).strip()
    if state != "device":
        raise RuntimeError(f"The default ADB device is unavailable: {state}")


def deploy_runtime(device_root: str, skip_upload: bool) -> str:
    remote_cli = f"{device_root}/genie_cli"
    if skip_upload:
        print("[1/4] Skipping runtime upload (--skip-upload).")
    else:
        print("[1/4] Uploading Genie CLI runtime ...")
        run_adb("shell", f"mkdir -p '{remote_cli}'")
        for relative in REQUIRED_FILES:
            local_path = LOCAL_ROOT / relative
            remote_path = f"{remote_cli}/{relative.as_posix()}"
            remote_parent = remote_path.rsplit("/", 1)[0]
            run_adb("shell", f"mkdir -p '{remote_parent}'")
            print(f"  push {relative.as_posix()}")
            run_adb("push", str(local_path), remote_path)

    print("[2/4] Applying device permissions ...")
    run_adb(
        "shell",
        f"chown -R 2000:2000 '{remote_cli}'; "
        f"chmod -R u+rwX,go+rX '{remote_cli}'; "
        f"chmod 755 '{remote_cli}/bin/genie-t2t-run'",
    )

    print("[3/4] Checking model and runtime files ...")
    checks = [
        ("genie_config.json", 31),
        ("part1_of_2.bin", 32),
        ("part2_of_2.bin", 33),
        ("tokenizer.json", 34),
        ("htp_backend_ext_config.json", 35),
        ("dsp/libQnnHtpV68Skel.so", 36),
        ("genie_cli/libs/libQnnHtp.so", 37),
        ("genie_cli/libs/libQnnSystem.so", 38),
        ("genie_cli/libs/libGenie.so", 41),
        ("genie_cli/libs/libQnnHtpNetRunExtensions.so", 42),
        ("genie_cli/libs/libQnnHtpV68Stub.so", 43),
    ]
    lines = [f"ROOT={device_root}"]
    for relative, code in checks:
        lines.append(
            f'test -f "$ROOT/{relative}" || '
            f'{{ echo "Missing {relative}"; exit {code}; }}'
        )
    lines.append(
        'test -x "$ROOT/genie_cli/bin/genie-t2t-run" || '
        '{ echo "genie-t2t-run is not executable"; exit 39; }'
    )
    shell_base64("\n".join(lines))
    return remote_cli



def make_raw_prompt(user_prompt: str) -> str:
    return (
        "<|im_start|>system\n"
        "English request classifier.\n"
        "Unsafe/illegal/harmful -> REJECT\n"
        "Supported vehicle control -> exact command\n"
        "Otherwise -> UNKNOWN\n"
        "Priority: REJECT > vehicle command > UNKNOWN.\n"
        "Output only the result.\n"
        "<|im_end|>\n<|im_start|>user\n"
        f"{user_prompt}\n"
        "<|im_end|>\n"
        "<|im_start|>assistant\n"
    )

def make_android_aligned_config(
    device_root: str, max_num_tokens: int
) -> dict[str, Any]:
    """Build the same Genie JSON emitted by the Android JNI defaults."""
    return {
        "dialog": {
            "version": 1,
            "type": "basic",
            "stop-sequence": ["<|im_end|>"],
            "max-num-tokens": max_num_tokens,
            "context": {
                "version": 1,
                "size": ANDROID_CONTEXT_SIZE,
                "n-vocab": 151936,
                "bos-token": 151643,
                "eos-token": 151645,
            },
            "sampler": {
                "version": 1,
                "seed": ANDROID_SEED,
                "temp": ANDROID_TEMPERATURE,
                "top-k": ANDROID_TOP_K,
                "top-p": ANDROID_TOP_P,
                "greedy": True,
                "token-penalty": {
                    "version": 1,
                    "penalize-last-n": 128,
                    "repetition-penalty": 1.0,
                    "presence-penalty": ANDROID_PRESENCE_PENALTY,
                    "frequency-penalty": 0.0,
                },
            },
            "tokenizer": {
                "version": 1,
                "path": f"{device_root}/tokenizer.json",
            },
            "engine": {
                "version": 1,
                "n-threads": ANDROID_THREAD_COUNT,
                "backend": {
                    "version": 1,
                    "type": "QnnHtp",
                    "QnnHtp": {
                        "version": 1,
                        "use-mmap": True,
                        "spill-fill-bufsize": 0,
                        "mmap-budget": 0,
                        "poll": False,
                        "cpu-mask": "0xe0",
                        "kv-dim": 128,
                        "allow-async-init": False,
                        "pos-id-dim": 64,
                        "rope-theta": 1000000,
                    },
                    "extensions": f"{device_root}/htp_backend_ext_config.json",
                },
                "model": {
                    "version": 1,
                    "type": "binary",
                    "binary": {
                        "version": 1,
                        "ctx-bins": [
                            f"{device_root}/part1_of_2.bin",
                            f"{device_root}/part2_of_2.bin",
                        ],
                    },
                },
            },
        }
    }


def print_prompt_tokens(
    prompt: str, tokenizer: Tokenizer, print_token_ids: bool
) -> int:
    user_encoding = tokenizer.encode(prompt, add_special_tokens=False)
    raw_encoding = tokenizer.encode(make_raw_prompt(prompt), add_special_tokens=False)
    genie_prompt_token_count = len(raw_encoding.ids) + ANDROID_BOS_TOKEN_COUNT
    print(f"  User prompt tokens     : {len(user_encoding.ids)}", flush=True)
    print(f"  Full raw prompt tokens : {len(raw_encoding.ids)}", flush=True)
    print(f"  Genie prompt tokens    : {genie_prompt_token_count}", flush=True)
    if print_token_ids:
        print(f"  User prompt token IDs  : {user_encoding.ids}", flush=True)
        print(f"  Full raw token IDs     : {raw_encoding.ids}", flush=True)
    return genie_prompt_token_count


def parse_metrics(profile_text: str) -> Metrics:
    profile = json.loads(profile_text)
    query_events: list[dict[str, Any]] = []
    for component in profile.get("components", []):
        query_events.extend(
            event
            for event in component.get("events", [])
            if event.get("type") == "GenieDialog_query"
        )
    if not query_events:
        raise ValueError("No GenieDialog_query event found in device profile")
    event = query_events[-1]
    prefill = float(event["time-to-first-token"]["value"]) / 1_000_000
    decode = float(event["token-generation-time"]["value"]) / 1_000_000
    return Metrics(
        prefill_seconds=prefill,
        decode_seconds=decode,
        total_seconds=prefill + decode,
        prefill_rate=float(event["prompt-processing-rate"]["value"]),
        decode_rate=float(event["token-generation-rate"]["value"]),
        prompt_tokens=int(event["num-prompt-tokens"]["value"]),
        decode_tokens=int(event["num-generated-tokens"]["value"]),
    )


def run_case(
    prompt: str,
    prompt_token_count: int,
    case_number: int,
    device_root: str,
    remote_cli: str,
    temporary_directory: Path,
) -> tuple[str, Metrics]:
    if prompt_token_count >= ANDROID_MAX_ALL_TOKENS:
        raise ValueError(
            f"Input token count {prompt_token_count} reaches Android max_all_token "
            f"{ANDROID_MAX_ALL_TOKENS}; no output-token budget remains."
        )
    output_token_budget = min(
        ANDROID_MAX_OUTPUT_TOKENS,
        ANDROID_MAX_ALL_TOKENS - prompt_token_count,
    )
    local_prompt = temporary_directory / f"prompt-{case_number:05d}.txt"
    local_prompt.write_text(make_raw_prompt(prompt), encoding="utf-8", newline="")
    local_config = temporary_directory / f"android-config-{case_number:05d}.json"
    local_config.write_text(
        json.dumps(
            make_android_aligned_config(device_root, output_token_budget),
            ensure_ascii=False,
            separators=(",", ":"),
        ),
        encoding="utf-8",
        newline="",
    )
    remote_prompt = f"{remote_cli}/batch_prompt.txt"
    remote_config = f"{remote_cli}/android_aligned_genie_config.json"
    remote_profile = f"{remote_cli}/batch_profile.json"
    run_adb("push", str(local_prompt), remote_prompt)
    run_adb("push", str(local_config), remote_config)

    script = f"""cd {device_root} || exit 40
export LD_LIBRARY_PATH=$PWD/genie_cli/libs:$PWD/dsp:/vendor/lib64:/system/lib64
export ADSP_LIBRARY_PATH=\"/vendor/lib/rfsa/adsp;$PWD;$PWD/dsp;$PWD/lib\"
export CDSP_LIBRARY_PATH=\"$ADSP_LIBRARY_PATH\"
export CDSP1_LIBRARY_PATH=\"$ADSP_LIBRARY_PATH\"
rm -f '{remote_profile}'
./genie_cli/bin/genie-t2t-run \\
  --config '{remote_config}' \\
  --prompt_file '{remote_prompt}' \\
  --log verbose \\
  --profile '{remote_profile}'
"""
    console_text = shell_base64(script, use_shell_user=True)
    match = re.search(r"\[BEGIN\]:\s*(.*?)\s*\[END\]", console_text, re.DOTALL)
    if not match:
        raise ValueError("Could not find [BEGIN]: ... [END] in genie-t2t-run output")
    raw_output = match.group(1).strip()
    # Extract only the JSON object: find first '{' and last '}'
    json_start = raw_output.find("{")
    json_end = raw_output.rfind("}")
    if json_start != -1 and json_end != -1 and json_end > json_start:
        actual = raw_output[json_start:json_end + 1]
    else:
        actual = raw_output
    profile_text = run_adb("shell", f"cat '{remote_profile}'", capture=True)
    return actual, parse_metrics(profile_text)


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def wrap_top(cell) -> None:
    alignment = copy(cell.alignment)
    alignment.wrap_text = True
    alignment.vertical = "top"
    cell.alignment = alignment


def read_cases(input_path: Path, max_cases: int) -> list[Case]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if header_values is None:
            raise ValueError(f"Input workbook is empty: {input_path}")
        header_map = {
            str(value).strip(): index
            for index, value in enumerate(header_values)
            if value is not None
        }
        missing = [header for header in HEADERS[:2] if header not in header_map]
        if missing:
            raise ValueError(f"Input workbook is missing columns: {', '.join(missing)}")
        cases = []
        for row_values in rows:
            prompt_value = row_values[header_map[HEADERS[0]]]
            expected_value = row_values[header_map[HEADERS[1]]]
            prompt = "" if prompt_value is None else str(prompt_value).strip()
            if prompt:
                cases.append(
                    Case(
                        prompt=prompt,
                        expected="" if expected_value is None else str(expected_value),
                    )
                )
    finally:
        workbook.close()
    if max_cases > 0:
        cases = cases[:max_cases]
    if not cases:
        raise ValueError(f"No non-empty prompts found in {input_path}")
    return cases


def prepare_result(
    reference_path: Path,
    output_path: Path,
    cases: list[Case],
    tokenizer: Tokenizer,
) -> None:
    if output_path.exists():
        raise FileExistsError(f"Output workbook already exists: {output_path}")
    shutil.copy2(reference_path, output_path)
    workbook = load_workbook(output_path)
    try:
        sheet = workbook.worksheets[0]
        sheet.title = "批量推理结果"
        # The reference workbook has 11 columns. Insert the new token-ID column
        # after 输入提示词 and 预期结果 so the original result/metric columns shift right.
        sheet.insert_cols(3, 1)
        for column, header in enumerate(HEADERS, start=1):
            sheet.cell(1, column).value = header
        sheet.cell(1, 3)._style = copy(sheet.cell(1, 4)._style)

        template = [copy(sheet.cell(2, column)._style) for column in range(1, 12)]
        template.append(copy(sheet.cell(2, 12)._style))
        template[2] = copy(sheet.cell(2, 4)._style)
        target_last_row = len(cases) + 1
        clear_last_row = max(sheet.max_row, target_last_row)
        for row in range(2, clear_last_row + 1):
            for column in range(1, 13):
                cell = sheet.cell(row, column)
                cell.value = None
                cell._style = copy(template[column - 1])
        for row, case in enumerate(cases, start=2):
            sheet.cell(row, 1).value = case.prompt
            sheet.cell(row, 2).value = case.expected
            input_token_ids = tokenizer.encode(
                make_raw_prompt(case.prompt), add_special_tokens=False
            ).ids
            sheet.cell(row, 3).value = json.dumps(
                input_token_ids, ensure_ascii=False, separators=(",", ":")
            )
            for column in range(1, 6):
                wrap_top(sheet.cell(row, column))
        if sheet.max_row > target_last_row:
            sheet.delete_rows(target_last_row + 1, sheet.max_row - target_last_row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:L{target_last_row}"
        sheet.column_dimensions["C"].width = 48
        workbook.save(output_path)
    finally:
        workbook.close()


def update_result(
    output_path: Path,
    row: int,
    actual: str,
    comparison: str,
    metrics: Metrics | None,
) -> None:
    workbook = load_workbook(output_path)
    try:
        sheet = workbook.worksheets[0]
        sheet.cell(row, 4).value = actual
        sheet.cell(row, 5).value = comparison
        if metrics is not None:
            values = list(asdict(metrics).values())
            for column, value in enumerate(values, start=6):
                sheet.cell(row, column).value = value
            for column in range(6, 11):
                sheet.cell(row, column).number_format = "0.000000"
            for column in range(11, 13):
                sheet.cell(row, column).number_format = "0"
        for column in range(1, 6):
            wrap_top(sheet.cell(row, column))
        workbook.save(output_path)
    finally:
        workbook.close()


def comparable_json(text: str) -> Any | None:
    candidate = text.strip()
    fence_match = re.fullmatch(
        r"```(?:json)?\s*(.*?)\s*```", candidate, re.DOTALL | re.IGNORECASE
    )
    if fence_match:
        candidate = fence_match.group(1).strip()
    try:
        value = json.loads(candidate)
    except (json.JSONDecodeError, TypeError):
        return None
    return value if isinstance(value, (dict, list)) else None


def json_values_equal(expected: Any, actual: Any) -> bool:
    if isinstance(expected, dict) and isinstance(actual, dict):
        return expected.keys() == actual.keys() and all(
            json_values_equal(expected[key], actual[key]) for key in expected
        )
    if isinstance(expected, list) and isinstance(actual, list):
        return len(expected) == len(actual) and all(
            json_values_equal(expected_item, actual_item)
            for expected_item, actual_item in zip(expected, actual)
        )
    if isinstance(expected, bool) or isinstance(actual, bool):
        return type(expected) is type(actual) and expected == actual
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return expected == actual
    return type(expected) is type(actual) and expected == actual


def compare_answers(expected: str, actual: str) -> bool:
    expected_text = "" if expected is None else expected.strip()
    actual_text = "" if actual is None else actual.strip()
    expected_json = comparable_json(expected_text)
    actual_json = comparable_json(actual_text)
    if expected_json is not None and actual_json is not None:
        return json_values_equal(expected_json, actual_json)
    return expected_text == actual_text


def batch_inference(
    args: argparse.Namespace, remote_cli: str, tokenizer: Tokenizer
) -> Path:
    input_path = Path(args.input).resolve()
    reference_path = Path(args.reference).resolve()
    result_directory = LOCAL_ROOT / "batch_result"
    result_directory.mkdir(parents=True, exist_ok=True)
    if args.output:
        requested_output = Path(args.output)
        output_path = (
            requested_output.resolve()
            if requested_output.is_absolute()
            else (result_directory / requested_output).resolve()
        )
    else:
        output_path = result_directory / (
            f"batch_result_Qwen3-0.6B-QNN-v17_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output workbook must use the .xlsx extension")
    cases = read_cases(input_path, args.max_cases)
    prepare_result(reference_path, output_path, cases, tokenizer)

    print(f"Batch input : {input_path}")
    print(f"Cases       : {len(cases)}")
    print(f"Output      : {output_path}")

    with tempfile.TemporaryDirectory(prefix="genie-batch-") as temporary:
        temporary_directory = Path(temporary)
        for index, case in enumerate(cases, start=1):
            print(f"\n[{index}/{len(cases)}] {case.prompt}")
            prompt_token_count = print_prompt_tokens(
                case.prompt, tokenizer, args.print_token_ids
            )
            try:
                actual, metrics = run_case(
                    case.prompt,
                    prompt_token_count,
                    index,
                    args.device_root,
                    remote_cli,
                    temporary_directory,
                )
                comparison = "一致" if compare_answers(case.expected, actual) else "不一致"
                update_result(output_path, index + 1, actual, comparison, metrics)
                print(
                    f"  Profile prompt tokens : {metrics.prompt_tokens}\n"
                    f"  {comparison}; total={metrics.total_seconds:.3f}s; result={actual}"
                )
            except Exception as error:  # keep later cases running and save the error row
                message = f"ERROR: {error}"
                update_result(output_path, index + 1, message, "错误", None)
                print(f"  {message}", file=sys.stderr)
    print(f"\nBatch inference completed.\nResult workbook: {output_path}")
    return output_path


def single_inference(
    args: argparse.Namespace, remote_cli: str, tokenizer: Tokenizer
) -> None:
    prompt_file = LOCAL_ROOT / "genie_cli_prompt.txt"
    prompt_text = prompt_file.read_text(encoding="utf-8")
    match = re.search(
        r"<\|im_start\|>user\s*(.*?)<\|im_end\|>", prompt_text, re.DOTALL
    )
    if not match:
        raise ValueError(f"Could not extract user prompt from {prompt_file}")
    user_prompt = match.group(1).strip()
    print(f"Prompt: {user_prompt}")
    prompt_token_count = print_prompt_tokens(
        user_prompt, tokenizer, args.print_token_ids
    )
    with tempfile.TemporaryDirectory(prefix="genie-single-") as temporary:
        actual, metrics = run_case(
            user_prompt,
            prompt_token_count,
            1,
            args.device_root,
            remote_cli,
            Path(temporary),
        )
    print(f"Result: {actual}")
    print(f"Total : {metrics.total_seconds:.3f}s")
    print(f"Profile prompt tokens: {metrics.prompt_tokens}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("batch", "single"), default="batch")
    parser.add_argument("--device-root", default=DEFAULT_DEVICE_ROOT)
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--reference", default=str(DEFAULT_REFERENCE))
    parser.add_argument("--output", default="")
    parser.add_argument("--tokenizer", default=str(DEFAULT_TOKENIZER))
    token_id_group = parser.add_mutually_exclusive_group()
    token_id_group.add_argument(
        "--print-token-ids",
        dest="print_token_ids",
        action="store_true",
        help="Print user and full raw prompt token IDs before inference (default).",
    )
    token_id_group.add_argument(
        "--no-print-token-ids",
        dest="print_token_ids",
        action="store_false",
        help="Do not print token ID lists.",
    )
    parser.set_defaults(print_token_ids=True)
    parser.add_argument("--max-cases", type=int, default=0)
    parser.add_argument("--skip-upload", action="store_true")
    parser.add_argument("--batch-engine", choices=("resident", "legacy"), default="resident")
    parser.add_argument("--recover-jsonl", default="", help="Export a saved resident journal without running inference")
    return parser


def main() -> int:
    configure_console()
    args = build_parser().parse_args()
    try:
        if args.recover_jsonl:
            from persistent_batch import export_journal
            _, errors = export_journal(Path(args.recover_jsonl), sys.modules[__name__])
            return 1 if errors else 0
        if args.mode == "batch" and args.batch_engine == "resident":
            if not (LOCAL_ROOT / "bin/genie-batch-run").is_file():
                raise FileNotFoundError("Missing bin/genie-batch-run; run native/build.ps1 first")
        validate_inputs(args)
        tokenizer = Tokenizer.from_file(str(Path(args.tokenizer).resolve()))
        remote_cli = deploy_runtime(args.device_root, args.skip_upload)
        print(f"[4/4] Starting {args.mode} inference ...")
        if args.mode == "batch":
            if args.batch_engine == "resident":
                from persistent_batch import run
                run(args, remote_cli, tokenizer, sys.modules[__name__])
            else:
                batch_inference(args, remote_cli, tokenizer)
        else:
            single_inference(args, remote_cli, tokenizer)
        return 0
    except Exception as error:
        print(f"\nFAILED: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
